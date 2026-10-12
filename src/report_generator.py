"""STEP 13 - Downloadable business report (DOCX + Markdown) and the trend chart."""
from __future__ import annotations

import io
from datetime import datetime
from typing import Dict, Optional

import matplotlib
matplotlib.use("Agg")  # no GUI needed (works on servers and in Streamlit)
import matplotlib.pyplot as plt
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from src.llm_generator import AIInsightReport
from matplotlib.ticker import FuncFormatter

from src.utils import fmt_pct, format_money, get_currency


def make_trend_chart(monthly: pd.DataFrame, months: int = 12) -> bytes:
    """Revenue (bars) and profit (line) for the last N months, as PNG bytes."""
    data = monthly.tail(months)
    labels = [str(m) for m in data["Month"]]
    fig, ax = plt.subplots(figsize=(8, 3.6))
    ax.bar(labels, data["Revenue"], color="#4C78A8", label="Revenue")
    ax.plot(labels, data["Profit"], color="#F58518", marker="o", label="Profit")
    ax.set_ylabel(f"Amount ({get_currency()})")
    ax.yaxis.set_major_formatter(FuncFormatter(lambda v, _: f"{v:,.0f}"))
    ax.set_title("Monthly Revenue and Profit")
    ax.tick_params(axis="x", rotation=45)
    ax.legend(frameon=False)
    ax.spines[["top", "right"]].set_visible(False)
    fig.tight_layout()
    buf = io.BytesIO()
    fig.savefig(buf, format="png", dpi=150)
    plt.close(fig)
    return buf.getvalue()


# order in which driver tables appear in the report (only dimensions present in the data are shown)
REPORT_DIMENSIONS = ("Region", "Product_Category", "Sub_Category", "Sales_Channel", "Payment_Method",
                     "Customer_Segment", "Product")


def _kpi_rows(payload: Dict):
    k = payload["kpis"]
    rows = [
        ("Revenue", k["revenue"]["current_display"], k["revenue"]["previous_display"], fmt_pct(k["revenue"]["growth_pct"])),
        ("Profit", k["profit"]["current_display"], k["profit"]["previous_display"], fmt_pct(k["profit"]["growth_pct"])),
        ("Orders", f"{k['orders']['current']:,}", f"{k['orders']['previous']:,}", fmt_pct(k["orders"]["growth_pct"])),
        ("Customers", f"{k['customers']['current']:,}", f"{k['customers']['previous']:,}", fmt_pct(k["customers"]["growth_pct"])),
        ("Average Order Value", k["average_order_value"]["current_display"], k["average_order_value"]["previous_display"], fmt_pct(k["average_order_value"]["growth_pct"])),
        ("Profit Margin", f"{k['profit_margin_pct']['current']}%", f"{k['profit_margin_pct']['previous']}%", f"{k['profit_margin_pct']['change_pp']:+.1f} pp"),
    ]
    return rows


def _driver_rows(payload: Dict, dimension: str):
    d = payload["drivers"].get(dimension, {})
    rows = []
    for tag, key in (("Driver", "top_positive"), ("Decliner", "top_negative")):
        for m in d.get(key, []):
            rows.append((tag, m["name"], fmt_pct(m["growth_pct"]), m["revenue_change_display"]))
    return rows


def _forecast_rows(payload: Dict):
    fc = payload.get("forecast", {})
    if not fc.get("available"):
        return []
    return [(p["month"], p["value_display"], f"{p['lower_display']} - {p['upper_display']}") for p in fc["projections"]]


def _forecast_note(payload: Dict) -> str:
    fc = payload.get("forecast", {})
    if not fc.get("available"):
        return "Forecast not available: " + str(fc.get("reason", "insufficient data")) + "."
    note = f"Method: {fc['method'].replace('_', ' ')}. Statistical projection with a 95% range, not a guarantee."
    if fc.get("backtest_mape_pct") is not None:
        note += f" Backtest error on the last 3 known months: {fc['backtest_mape_pct']}% (MAPE)."
    return note


def _anomaly_lines(payload: Dict):
    an = payload.get("anomalies", {"daily": [], "members": []})
    lines = [f"{a['name']} ({a['dimension'].replace('_', ' ')}): MoM {a['mom_pct']:+.1f}% vs typical {a['typical_mom_pct']:+.1f}% (z={a['z_score']}, {a['direction']})"
             for a in an["members"]]
    lines += [f"{a['date']}: revenue {a['revenue_display']} vs typical {a['typical_display']} ({a['direction']}, robust z={a['robust_z_score']})"
              for a in an["daily"]]
    return lines


def _trend_sentence(payload: Dict) -> str:
    t = payload["trend"]["revenue"]
    parts = [f"Revenue trend over the last months is {t['direction']}"]
    if t["slope_pct_per_month"] is not None:
        parts[0] += f" ({t['slope_pct_per_month']:+.1f}% per month on average)"
    if t["latest_yoy_pct"] is not None:
        parts.append(f"year-over-year change for the latest month is {t['latest_yoy_pct']:+.1f}%")
    return "; ".join(parts) + "."


def build_markdown(payload: Dict, report: AIInsightReport, meta: Optional[Dict] = None) -> str:
    meta = meta or {}
    lines = [f"# Business Insight Report - {payload['period']['label']}", "",
             f"_Generated on {datetime.now():%d %b %Y %H:%M}. Source: "
             f"{'AI-generated (' + str(meta.get('model')) + ')' if meta.get('source') == 'llm' else 'rule-based engine'}._", ""]
    for w in meta.get("warnings", []):
        lines += [f"> Note: {w}", ""]
    lines += ["## Executive Summary", report.executive_summary, "", "## KPI Summary",
              "| Metric | Current | Previous | Change |", "|---|---|---|---|"]
    lines += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in _kpi_rows(payload)]
    lines += ["", "## Trend Analysis", _trend_sentence(payload)]
    frows = _forecast_rows(payload)
    lines += ["", "## Forecast"]
    if frows:
        lines += ["| Month | Projected revenue | 95% range |", "|---|---|---|"] + [f"| {a} | {b} | {c} |" for a, b, c in frows]
    lines += ["", _forecast_note(payload)]
    anomalies = _anomaly_lines(payload)
    lines += ["", "## Unusual Movements (anomalies)"] + ([f"- {x}" for x in anomalies] or ["- None detected."])
    lines += ["", "## Growth Drivers"]
    lines += [f"- {x}" for x in report.key_growth_drivers] or ["- None identified."]
    for dim in REPORT_DIMENSIONS:
        rows = _driver_rows(payload, dim)
        if rows:
            lines += ["", f"**By {dim.replace('_', ' ')}**", "| Type | Name | Growth | Revenue change |", "|---|---|---|---|"]
            lines += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in rows]
    lines += ["", "## Negative Contributors"] + ([f"- {x}" for x in report.negative_contributors] or ["- None identified."])
    lines += ["", "## Risks / Areas to Investigate"] + ([f"- {x}" for x in report.risks_to_investigate] or ["- None identified."])
    lines += ["", "## Recommended Actions (suggestions, not facts)"] + ([f"- {x}" for x in report.recommended_actions] or ["- None."])
    lines += ["", "---", "_Method: all figures are computed in Python from the cleaned dataset; the AI only phrases verified results._"]
    return "\n".join(lines)


def _bullets(doc: Document, items, empty: str = "None identified."):
    for text in items or [empty]:
        doc.add_paragraph(text, style="List Bullet")


def _table(doc: Document, header, rows):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Light Grid Accent 1"
    for i, h in enumerate(header):
        t.rows[0].cells[i].text = h
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v)
    doc.add_paragraph()


def build_docx(payload: Dict, report: AIInsightReport, meta: Optional[Dict] = None,
               monthly: Optional[pd.DataFrame] = None, cleaning_report: Optional[Dict] = None) -> bytes:
    meta = meta or {}
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    title = doc.add_heading("Business Insight Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    doc.add_paragraph(f"Period: {payload['period']['label']}  |  Generated: {datetime.now():%d %b %Y %H:%M}  |  "
                      f"Source: {'AI (' + str(meta.get('model')) + ')' if meta.get('source') == 'llm' else 'rule-based engine'}")
    for w in meta.get("warnings", []):
        doc.add_paragraph(f"Note: {w}").runs[0].italic = True

    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(report.executive_summary)

    doc.add_heading("KPI Summary", level=1)
    _table(doc, ["Metric", "Current", "Previous", "Change"], _kpi_rows(payload))

    doc.add_heading("Trend Analysis", level=1)
    doc.add_paragraph(_trend_sentence(payload))
    if monthly is not None and len(monthly) >= 2:
        doc.add_picture(io.BytesIO(make_trend_chart(monthly)), width=Inches(6.2))

    doc.add_heading("Forecast", level=1)
    frows = _forecast_rows(payload)
    if frows:
        _table(doc, ["Month", "Projected revenue", "95% range"], frows)
    doc.add_paragraph(_forecast_note(payload)).runs[0].italic = True
    doc.add_heading("Unusual Movements (anomalies)", level=1)
    _bullets(doc, _anomaly_lines(payload), "None detected.")

    doc.add_heading("Growth Drivers", level=1)
    _bullets(doc, report.key_growth_drivers)
    for dim in REPORT_DIMENSIONS:
        rows = _driver_rows(payload, dim)
        if rows:
            doc.add_paragraph(f"By {dim.replace('_', ' ')}").runs[0].bold = True
            _table(doc, ["Type", "Name", "Growth", "Revenue change"], rows)

    doc.add_heading("Negative Contributors", level=1)
    _bullets(doc, report.negative_contributors)
    doc.add_heading("Risks / Areas to Investigate", level=1)
    _bullets(doc, report.risks_to_investigate)
    doc.add_heading("Recommended Actions", level=1)
    doc.add_paragraph("These are suggestions for the business to consider, not observed facts.").runs[0].italic = True
    _bullets(doc, report.recommended_actions, "None.")

    if cleaning_report:
        doc.add_heading("Appendix: Data Quality", level=1)
        doc.add_paragraph(f"Rows received: {cleaning_report['rows_in']:,}. Rows used after cleaning: {cleaning_report['rows_out']:,}.")
        _bullets(doc, [f"{k.replace('_', ' ')}: {v}" for k, v in cleaning_report["issues"].items()], "No issues found.")
    doc.add_paragraph("Method: all figures are computed in Python from the cleaned dataset; the AI only phrases verified results.").runs[0].italic = True

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
